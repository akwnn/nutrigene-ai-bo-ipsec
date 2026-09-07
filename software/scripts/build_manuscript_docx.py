#!/usr/bin/env python3
"""Build the canonical SPADE review manuscript as a reproducible Word file."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_SOURCE = ROOT / "publication" / "manuscript" / "MANUSCRIPT.md"
DEFAULT_CITATION = ROOT / "CITATION.cff"
DEFAULT_OUTPUT = ROOT / "publication" / "manuscript" / "SPADE-MANUSCRIPT.docx"
CONTENT_WIDTH_DXA = 9360
BASE_FONT = "Times New Roman"
INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|(?<!\*)\*[^*]+?\*(?!\*))")


def set_run_font(run, *, name: str = BASE_FONT, size: float = 12) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), name)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction, separate, value, end):
        run._r.append(node)
    set_run_font(run, size=10)


def configure_footer(footer) -> None:
    paragraph = footer.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    set_run_font(run, size=10)
    add_page_field(paragraph)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1.5)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

    line_numbers = OxmlElement("w:lnNumType")
    line_numbers.set(qn("w:countBy"), "1")
    line_numbers.set(qn("w:start"), "1")
    line_numbers.set(qn("w:restart"), "continuous")
    line_numbers.set(qn("w:distance"), "360")
    section._sectPr.append(line_numbers)

    document.settings.odd_and_even_pages_header_footer = False
    section.different_first_page_header_footer = False
    for footer in (section.footer, section.even_page_footer, section.first_page_footer):
        configure_footer(footer)

    normal = document.styles["Normal"]
    normal.font.name = BASE_FONT
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:ascii"), BASE_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), BASE_FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    normal.paragraph_format.widow_control = True

    title = document.styles["Title"]
    title.font.name = BASE_FONT
    title.font.size = Pt(14)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title._element.rPr.rFonts.set(qn("w:ascii"), BASE_FONT)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), BASE_FONT)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    title.paragraph_format.keep_with_next = True

    heading_settings = {
        "Heading 1": (12, True, False, 12, 6),
        "Heading 2": (12, True, True, 10, 4),
    }
    for name, (size, bold, italic, before, after) in heading_settings.items():
        style = document.styles[name]
        style.font.name = BASE_FONT
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.italic = italic
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:ascii"), BASE_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), BASE_FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True


def add_inline_markdown(paragraph, text: str, *, size: float = 12) -> None:
    """Add the small inline Markdown subset used by the canonical manuscript."""
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_run_font(run, size=size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(token[1:-1])
        set_run_font(run, size=size)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size)


def author_names(citation_path: Path) -> list[str]:
    """Read author names from the simple CFF author list without a YAML dependency."""
    text = citation_path.read_text(encoding="utf-8")
    match = re.search(r"(?ms)^authors:\s*\n(?P<body>.*?)(?=^[A-Za-z][\w-]*:|\Z)", text)
    if match is None:
        raise ValueError(f"no authors block found in {citation_path}")
    names: list[str] = []
    entries = re.split(r"(?m)^\s*-\s+", match.group("body"))
    for entry in entries:
        family_match = re.search(r"(?m)^\s*family-names:\s*[\"']?([^\n\"']+)", entry)
        given_match = re.search(r"(?m)^\s*given-names:\s*[\"']?([^\n\"']+)", entry)
        if family_match is None and given_match is None:
            continue
        given = given_match.group(1).strip() if given_match else ""
        family = family_match.group(1).strip() if family_match else ""
        name = " ".join(part for part in (given, family) if part)
        if name:
            names.append(name)
    if not names:
        raise ValueError(f"no authors found in {citation_path}")
    return names


def join_names(names: list[str]) -> str:
    if len(names) == 1:
        return names[0]
    if len(names) == 2:
        return f"{names[0]} and {names[1]}"
    return f"{', '.join(names[:-1])}, and {names[-1]}"


def set_cell_margins(cell, *, top: int = 60, start: int = 70, bottom: int = 60, end: int = 70) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, *, color: str = "B7B7B7", size: str = "4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_row_contract(row, *, is_header: bool) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)
    if is_header:
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        tr_pr.append(repeat)


def table_widths(column_count: int) -> list[int]:
    widths = {
        3: [2500, 1900, 4960],
        4: [700, 3000, 2600, 3060],
        6: [850, 2250, 1000, 1550, 1200, 2510],
    }
    if column_count in widths:
        return widths[column_count]
    base, remainder = divmod(CONTENT_WIDTH_DXA, column_count)
    return [base + (1 if index < remainder else 0) for index in range(column_count)]


def add_markdown_table(document: Document, rows: list[list[str]]) -> None:
    if not rows or not rows[0]:
        return
    column_count = len(rows[0])
    if any(len(row) != column_count for row in rows):
        raise ValueError("Markdown table contains inconsistent column counts")

    table = document.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    widths = table_widths(column_count)
    font_size = 7.2 if column_count == 6 else 8.2
    horizontal_margin = 45 if column_count == 6 else 70
    grid_columns = table._tbl.tblGrid.gridCol_lst
    for grid_column, width in zip(grid_columns, widths):
        grid_column.set(qn("w:w"), str(width))

    for row_index, (word_row, values) in enumerate(zip(table.rows, rows)):
        is_header = row_index == 0
        set_row_contract(word_row, is_header=is_header)
        for cell, value, width in zip(word_row.cells, values, widths):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_width = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_width.set(qn("w:w"), str(width))
            tc_width.set(qn("w:type"), "dxa")
            set_cell_margins(cell, start=horizontal_margin, end=horizontal_margin)
            set_cell_border(cell)
            if is_header:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "D9EAF7")
                cell._tc.get_or_add_tcPr().append(shading)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            paragraph.paragraph_format.keep_together = True
            add_inline_markdown(paragraph, value, size=font_size)
            if is_header:
                for run in paragraph.runs:
                    run.bold = True

    document.add_paragraph().paragraph_format.space_after = Pt(0)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    raw_rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        raw_rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
        index += 1
    if len(raw_rows) < 2 or not all(re.fullmatch(r":?-{3,}:?", cell) for cell in raw_rows[1]):
        raise ValueError(f"malformed Markdown table beginning at line {start + 1}")
    return [raw_rows[0], *raw_rows[2:]], index


def add_authors(document: Document, names: list[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.space_after = Pt(18)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(join_names(names))
    set_run_font(run, size=12)


def add_body_paragraph(document: Document, text: str, *, in_references: bool) -> None:
    if text.startswith("- [ ] "):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.2)
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.space_after = Pt(2)
        add_inline_markdown(paragraph, f"\u2610 {text[6:]}", size=11)
        return

    paragraph = document.add_paragraph()
    if in_references and re.match(r"^\d+\. ", text):
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        paragraph.paragraph_format.space_after = Pt(4)
        add_inline_markdown(paragraph, text, size=10)
    else:
        add_inline_markdown(paragraph, text)


def build_document(source_path: Path, citation_path: Path, output_path: Path) -> None:
    lines = source_path.read_text(encoding="utf-8").splitlines()
    names = author_names(citation_path)
    document = Document()
    configure_document(document)

    document.core_properties.title = lines[0].removeprefix("# ").strip()
    document.core_properties.subject = "Fixed-well evaluation of conservative Gaussian-process design"
    document.core_properties.author = "; ".join(names)
    document.core_properties.keywords = "Bayesian optimization; Gaussian process; operating region; cell manufacturing"
    document.core_properties.comments = "Generated reproducibly from publication/manuscript/MANUSCRIPT.md"

    index = 0
    in_references = False
    title_added = False
    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            add_markdown_table(document, rows)
            continue
        if stripped.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            add_inline_markdown(paragraph, stripped[2:].strip(), size=14)
            add_authors(document, names)
            title_added = True
        elif stripped.startswith("## "):
            heading = stripped[3:].strip()
            paragraph = document.add_paragraph(style="Heading 1")
            if heading == "References":
                paragraph.paragraph_format.page_break_before = True
                in_references = True
            add_inline_markdown(paragraph, heading, size=12)
        elif stripped.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2")
            add_inline_markdown(paragraph, stripped[4:].strip(), size=12)
        elif stripped.startswith("**Table ") and stripped.endswith("**"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            paragraph.paragraph_format.keep_with_next = True
            add_inline_markdown(paragraph, stripped, size=10)
        else:
            add_body_paragraph(document, stripped, in_references=in_references)
        index += 1

    if not title_added:
        raise ValueError(f"canonical manuscript has no level-one title: {source_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    print(f"built manuscript: {output_path}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--citation", type=Path, default=DEFAULT_CITATION)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    build_document(args.source.resolve(), args.citation.resolve(), args.output.resolve())


if __name__ == "__main__":
    main()
