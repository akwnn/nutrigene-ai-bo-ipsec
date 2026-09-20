#!/usr/bin/env python3
"""Build the PLOS ONE manuscript DOCX from its tracked Markdown source."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

import mathml2omml
from latex2mathml.converter import convert as latex_to_mathml
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor


CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}
BASE_FONT = "Times New Roman"
MATH_FONT = "Cambria Math"


def set_run_font(run, name: str = BASE_FONT, size: float = 12) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), name)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in CELL_MARGIN_DXA.items():
        tag = tc_mar.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            tc_mar.append(tag)
        tag.set(qn("w:w"), str(value))
        tag.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {CONTENT_WIDTH_DXA}: {widths}")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "D9D9D9")
        borders.append(border)
    tbl_pr.append(borders)

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_field(paragraph, field: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)
    set_run_font(run, size=10)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1.5)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    sect_pr = section._sectPr
    line_num = sect_pr.find(qn("w:lnNumType"))
    if line_num is None:
        line_num = OxmlElement("w:lnNumType")
        sect_pr.append(line_num)
    line_num.set(qn("w:countBy"), "1")
    line_num.set(qn("w:start"), "1")
    line_num.set(qn("w:restart"), "continuous")

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    set_run_font(run, size=10)
    add_field(paragraph, "PAGE")

    styles = document.styles
    title = styles["Title"]
    title.font.color.rgb = RGBColor(0, 0, 0)
    for border in list(title._element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    normal = styles["Normal"]
    normal.font.name = BASE_FONT
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:ascii"), BASE_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), BASE_FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    normal.paragraph_format.widow_control = True

    heading_tokens = {
        "Heading 1": (12, True, False, 12, 6),
        "Heading 2": (12, True, False, 10, 4),
        "Heading 3": (12, True, True, 8, 3),
    }
    for name, (size, bold, italic, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = BASE_FONT
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.italic = italic
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:ascii"), BASE_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), BASE_FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True


INLINE_RE = re.compile(
    r"(\$\$.+?\$\$|\$.+?\$|\*\*.+?\*\*|`.+?`|(?<!\*)\*[^*]+?\*(?!\*))"
)


def math_to_omml(latex: str, *, display: bool = False):
    """Convert a LaTeX fragment into an editable native Word equation."""
    # latex2mathml expands \mathrm{Var} into three adjacent identifiers. The
    # downstream OMML converter then gives Word licence to space those letters
    # like multiplied variables. An operator token preserves the intended
    # upright label as one editable run (also for EC, IC, rel, and straddle).
    latex = re.sub(r"\\mathrm\{([^{}]+)\}", r"\\operatorname{\1}", latex)
    mathml = latex_to_mathml(latex, display="block" if display else "inline")
    omml = mathml2omml.convert(mathml)
    omml = omml.replace("<m:oMath>", f"<m:oMath {nsdecls('m')}>", 1)
    root = parse_xml(omml)
    # mathml2omml omits the mandatory degree nodes for square roots. Word may
    # repair that omission, but LibreOffice renders it as empty placeholder
    # boxes, so complete the native OMML structure deterministically.
    for radical in root.iter(qn("m:rad")):
        if radical.find(qn("m:radPr")) is None:
            properties = OxmlElement("m:radPr")
            degree_hidden = OxmlElement("m:degHide")
            degree_hidden.set(qn("m:val"), "on")
            properties.append(degree_hidden)
            radical.insert(0, properties)
        if radical.find(qn("m:deg")) is None:
            radical.insert(1, OxmlElement("m:deg"))
    return root


def add_inline(paragraph, text: str, size: float = 12) -> None:
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_run_font(run, size=size)
        token = match.group(0)
        if token.startswith("$$") and token.endswith("$$"):
            paragraph._p.append(math_to_omml(token[2:-2], display=True))
        elif token.startswith("$") and token.endswith("$"):
            paragraph._p.append(math_to_omml(token[1:-1]))
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            set_run_font(run, size=size)
        elif token.startswith("`"):
            inner = token[1:-1]
            run = paragraph.add_run(inner)
            # Paths/filenames stay monospaced; short identifiers use body font
            # (italic) so the manuscript does not look like source code.
            if ("/" in inner) or inner.endswith((".md", ".py", ".json", ".docx", ".txt")):
                set_run_font(run, name="Courier New", size=max(size - 1, 9))
            else:
                run.italic = True
                set_run_font(run, size=size)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            set_run_font(run, size=size)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size)


def add_paragraph(document: Document, text: str, *, style: str | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    add_inline(paragraph, text)
    paragraph.paragraph_format.widow_control = True
    if re.match(r"^\*\*Table \d+\.", text):
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.keep_together = True


def add_title(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(12)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    run.bold = True
    set_run_font(run, size=14)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    idx = start
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[idx].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
        idx += 1
    return rows, idx


def column_widths(headers: list[str]) -> list[int]:
    count = len(headers)
    joined = " ".join(headers).lower()
    if count == 5 and "claim gate" in joined:
        return [1600, 1900, 2200, 1860, 1800]
    if count == 5 and "design" in joined and "primary purpose" in joined:
        return [2200, 3000, 760, 900, 2500]
    if count == 5:
        return [3000, 1050, 1050, 2100, 2160]
    if count == 4:
        return [3150, 1550, 1850, 2810]
    if count == 3:
        return [2200, 3580, 3580]
    base = CONTENT_WIDTH_DXA // count
    widths = [base] * count
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    count = len(rows[0])
    table = document.add_table(rows=len(rows), cols=count)
    table.style = "Table Grid"
    table.allow_autofit = False
    for row_idx, values in enumerate(rows):
        tr_pr = table.rows[row_idx]._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        for col_idx, value in enumerate(values):
            cell = table.cell(row_idx, col_idx)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            add_inline(paragraph, value, size=8.5)
            if row_idx == 0:
                for run in paragraph.runs:
                    run.bold = True
                set_cell_shading(cell, "E7E6E6")
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, column_widths(rows[0]))
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_equation(document: Document, equation: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph._p.append(math_to_omml(equation, display=True))


def add_figure(document: Document, markdown_path: Path, target: str) -> None:
    image_path = (markdown_path.parent / target).resolve()
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(6.45))


def add_caption(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.keep_together = True
    add_inline(paragraph, text, size=10)


def unwrap_hard_wrapped_prose(lines: list[str]) -> list[str]:
    """Join hard-wrapped prose lines so Word gets one paragraph per blank-line break."""
    def special(s: str) -> bool:
        s = s.strip()
        if not s:
            return True
        if s.startswith(("#", "|", "![", "$$", "**Fig ", "**Table ", "**S", ">", "```")):
            return True
        if s in ("\\[", "\\]"):
            return True
        if re.match(r"^[-*]\s", s) or re.match(r"^\d+\.\s", s):
            return True
        return False

    out: list[str] = []
    buf: list[str] = []
    fence = False
    eq = False

    def flush() -> None:
        nonlocal buf
        if buf:
            out.append(" ".join(x.strip() for x in buf))
            buf = []

    for line in lines:
        s = line.strip()
        if s.startswith("```"):
            flush()
            fence = not fence
            out.append(line)
            continue
        if fence:
            out.append(line)
            continue
        if s == "\\[":
            flush()
            eq = True
            out.append(line)
            continue
        if eq:
            out.append(line)
            if s == "\\]":
                eq = False
            continue
        if special(line):
            flush()
            out.append(line)
            continue
        buf.append(line)
    flush()
    return out



def build(markdown_path: Path, output_path: Path, *, include_figures: bool = True,
          cover_letter: bool = False) -> None:
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    lines = unwrap_hard_wrapped_prose(lines)
    document = Document()
    configure_document(document)
    if cover_letter:
        section = document.sections[0]
        section.bottom_margin = Inches(1)
        section._sectPr.remove(section._sectPr.find(qn("w:lnNumType")))
        for paragraph in section.footer.paragraphs:
            paragraph.clear()
        normal = document.styles["Normal"].paragraph_format
        normal.line_spacing_rule = WD_LINE_SPACING.SINGLE
        normal.space_after = Pt(8)

    idx = 0
    in_equation = False
    equation_lines: list[str] = []
    while idx < len(lines):
        stripped = lines[idx].strip()
        if not stripped:
            idx += 1
            continue
        if stripped == "\\[":
            in_equation = True
            equation_lines = []
            idx += 1
            continue
        if in_equation:
            if stripped == "\\]":
                add_equation(document, " ".join(equation_lines))
                in_equation = False
            else:
                equation_lines.append(stripped)
            idx += 1
            continue
        if idx == 0 and stripped.startswith("# "):
            add_title(document, stripped[2:].strip())
            idx += 1
            continue
        if stripped.startswith("### "):
            add_paragraph(document, stripped[4:].strip(), style="Heading 2")
            idx += 1
            continue
        if stripped.startswith("## "):
            heading = stripped[3:].strip()
            if heading in {"Abstract", "References", "Supporting information captions"}:
                document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            add_paragraph(document, heading, style="Heading 1")
            idx += 1
            continue
        image_match = re.fullmatch(r"!\[[^]]*\]\(([^)]+)\)", stripped)
        if image_match:
            if include_figures:
                add_figure(document, markdown_path, image_match.group(1))
            idx += 1
            continue
        if stripped.startswith("|"):
            rows, idx = parse_table(lines, idx)
            add_table(document, rows)
            continue
        if stripped.startswith("**Fig "):
            add_caption(document, stripped)
            idx += 1
            continue
        if re.match(r"^\d+\. ", stripped):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.first_line_indent = Inches(-0.25)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            paragraph.paragraph_format.space_after = Pt(4)
            add_inline(paragraph, stripped, size=10)
            idx += 1
            continue
        add_paragraph(document, stripped)
        idx += 1

    core = document.core_properties
    core.title = lines[0].removeprefix("# ")
    core.subject = "PLOS ONE cover letter" if cover_letter else "PLOS ONE research article manuscript"
    core.keywords = "SPADE; Bayesian optimization; response-surface methodology; design space"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    image_targets = re.findall(r"!\[[^]]*\]\(([^)]+)\)", "\n".join(lines)) if include_figures else []
    manifest = {
        "source_sha256": hashlib.sha256(markdown_path.read_bytes()).hexdigest(),
        "builder_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        "output_sha256": hashlib.sha256(output_path.read_bytes()).hexdigest(),
        "include_figures": include_figures,
        "cover_letter": cover_letter,
        "image_sha256": {
            target: hashlib.sha256((markdown_path.parent / target).read_bytes()).hexdigest()
            for target in image_targets
        },
    }
    output_path.with_suffix(".docx.manifest.json").write_text(
        json.dumps(manifest, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=Path("manuscript/SPADE-PLOS-ONE.md"))
    parser.add_argument("--output", type=Path, default=Path("manuscript/SPADE-PLOS-ONE.docx"))
    parser.add_argument("--submission", action="store_true",
                        help="Keep figure captions but omit images for separate PLOS uploads")
    parser.add_argument("--cover-letter", action="store_true",
                        help="Use single-spaced letter layout without manuscript line numbers")
    args = parser.parse_args()
    build(args.input.resolve(), args.output.resolve(), include_figures=not args.submission,
          cover_letter=args.cover_letter)


if __name__ == "__main__":
    main()
